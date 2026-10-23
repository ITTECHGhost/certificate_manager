# =============================================================================
# services/excel_import_service.py — Excel Import & Template Engine
# =============================================================================

import io
import re
import logging
from datetime import datetime
from typing import List, Dict, Tuple, Any, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

ARABIC_INDIC_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


def normalize_digits(val: Any) -> str:
    """Normalize Arabic/Indic numerals to standard ASCII digits."""
    if val is None:
        return ""
    return str(val).translate(ARABIC_INDIC_DIGITS).strip()


def parse_float_safe(val: Any) -> Optional[float]:
    """Safely parse float score from cell value (handling string numbers, decimals, commas)."""
    if val is None:
        return None
    cleaned = normalize_digits(val).replace(",", ".").strip()
    if not cleaned:
        return None
    try:
        num = float(cleaned)
        return num
    except (ValueError, TypeError):
        return None


# -----------------------------------------------------------------------------
# Template 1: Students Roster
# -----------------------------------------------------------------------------

STUDENT_ROSTER_HEADERS = [
    ("الرقم الوطني / الهوية / الجواز", 25),
    ("الاسم الكامل بالعربية (الرباعي واللقب)", 35),
    ("الاسم الكامل بالإنكليزية", 32),
    ("الجنس (ذكر / أنثى)", 16),
    ("تاريخ الميلاد (YYYY-MM-DD)", 22),
    ("المحافظة / مسقط الرأس", 22),
    ("الجنسية", 18),
    ("سنة القبول (مثال: 2023-2024)", 25),
    ("نوع الدراسة (صباحي / مسائي)", 22),
]

SAMPLE_STUDENTS = [
    ["199912345678", "علي حسن محمد جاسم", "Ali Hassan Mohammed Jasim", "ذكر", "2001-05-14", "بغداد", "عراقي", "2023-2024", "صباحي"],
    ["200087654321", "فاطمة أحمد حسين خضير", "Fatima Ahmed Hussein Khudhair", "أنثى", "2002-09-22", "البصرة", "عراقي", "2023-2024", "صباحي"],
    ["199854321098", "عمر خالد طارق سعيد", "Omar Khalid Tariq Saeed", "ذكر", "2000-11-03", "نينوى", "عراقي", "2023-2024", "مسائي"],
]

GOVERNORATES_LIST = [
    "بغداد", "البصرة", "نينوى", "أربيل", "كركوك", "النجف", "كربلاء",
    "بابل", "الأنبار", "ديالى", "صلاح الدين", "واسط", "ذي قار",
    "المثنى", "القادسية", "ميسان", "السليمانية", "دهوك"
]


def generate_students_template() -> bytes:
    """Generate styled RTL Excel template for bulk student demographic ingestion."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "بيانات الطلاب"
    ws.views.sheetView[0].rightToLeft = True

    # Styling definitions
    title_font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    sub_font = Font(name="Segoe UI", size=10, italic=True, color="CBD5E1")
    sub_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")

    data_font = Font(name="Segoe UI", size=11)
    sample_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # Row 1: Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(STUDENT_ROSTER_HEADERS))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = "استمارة إضافة بيانات الطلاب — نظام إدارة الشهادات والوثائق الأكاديمية"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 36

    # Row 2: Instructions
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(STUDENT_ROSTER_HEADERS))
    sub_cell = ws.cell(row=2, column=1)
    sub_cell.value = "ملاحظة: يرجى عدم تعديل عناوين الأعمدة في الصف 3. يمكن مسح الصفوف النموذجية وإضافة بيانات الطلبة الفعلية."
    sub_cell.font = sub_font
    sub_cell.fill = sub_fill
    sub_cell.alignment = center_align
    ws.row_dimensions[2].height = 24

    # Row 3: Headers
    ws.row_dimensions[3].height = 32
    for col_idx, (header_text, width) in enumerate(STUDENT_ROSTER_HEADERS, start=1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = header_text
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(width, len(header_text) + 4)

    # Rows 4..6: Sample Rows
    for row_offset, sample in enumerate(SAMPLE_STUDENTS, start=4):
        ws.row_dimensions[row_offset].height = 26
        for col_idx, val in enumerate(sample, start=1):
            cell = ws.cell(row=row_offset, column=col_idx)
            cell.value = val
            cell.font = data_font
            cell.fill = sample_fill
            cell.border = thin_border
            cell.alignment = right_align if col_idx in [2, 3] else center_align

    # Add Dropdown Validations
    dv_gender = DataValidation(type="list", formula1='"ذكر,أنثى"', allow_blank=True)
    ws.add_data_validation(dv_gender)
    dv_gender.add("D4:D500")

    dv_study = DataValidation(type="list", formula1='"صباحي,مسائي"', allow_blank=True)
    ws.add_data_validation(dv_study)
    dv_study.add("I4:I500")

    govs_formula = f'"{",".join(GOVERNORATES_LIST[:10])}"'
    dv_gov = DataValidation(type="list", formula1=govs_formula, allow_blank=True)
    ws.add_data_validation(dv_gov)
    dv_gov.add("F4:F500")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def parse_students_file(file_bytes: bytes) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Parse uploaded Students Roster Excel file.
    Returns: (parsed_students_list, error_messages_list)
    """
    students: List[Dict[str, Any]] = []
    errors: List[str] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
    except Exception as exc:
        return [], [f"تعذر فتح ملف الـ Excel: {exc}"]

    # Locate header row (search for 'الاسم' or 'الرقم الوطني')
    header_row_idx = None
    col_mapping: Dict[str, int] = {}

    for r in range(1, min(10, ws.max_row + 1)):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r, column=c).value or "").strip()
            if "الاسم" in val or "الرقم الوطني" in val:
                header_row_idx = r
                break
        if header_row_idx:
            break

    if not header_row_idx:
        return [], ["لم يتم العثور على عناوين الأعمدة في ملف Excel. يرجى استخدام النموذج الرسمي."]

    # Build column mapping
    for c in range(1, ws.max_column + 1):
        raw_val = str(ws.cell(row=header_row_idx, column=c).value or "").strip()
        if not raw_val:
            continue
        if "الرقم الوطني" in raw_val or "الهوية" in raw_val or "الجواز" in raw_val:
            col_mapping["national_id"] = c
        elif "الاسم" in raw_val and ("عربي" in raw_val or "الرباعي" in raw_val or "العربية" in raw_val):
            col_mapping["full_name_ar"] = c
        elif "الاسم" in raw_val and ("إنكليزي" in raw_val or "انكليزي" in raw_val or "English" in raw_val):
            col_mapping["full_name_en"] = c
        elif "الجنس" in raw_val:
            col_mapping["gender"] = c
        elif "الميلاد" in raw_val or "تاريخ" in raw_val:
            col_mapping["date_of_birth"] = c
        elif "المحافظة" in raw_val or "مسقط" in raw_val or "الولادة" in raw_val:
            col_mapping["birthplace"] = c
        elif "الجنسية" in raw_val:
            col_mapping["nationality"] = c
        elif "القبول" in raw_val or "سنة" in raw_val:
            col_mapping["admission_year"] = c
        elif "نوع الدراسة" in raw_val or "الدراسة" in raw_val:
            col_mapping["study_type"] = c

    # Fallback to standard column indexes if some headers weren't named identically
    if "full_name_ar" not in col_mapping:
        col_mapping["full_name_ar"] = 2
    if "national_id" not in col_mapping:
        col_mapping["national_id"] = 1

    seen_ids = set()
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        name_ar_val = str(ws.cell(row=row_idx, column=col_mapping.get("full_name_ar", 2)).value or "").strip()
        if not name_ar_val:
            continue  # skip empty row

        nid_val = normalize_digits(ws.cell(row=row_idx, column=col_mapping.get("national_id", 1)).value)
        if not nid_val:
            # Generate fallback temp identifier if missing
            nid_val = f"GEN_{row_idx}_{int(datetime.now().timestamp()) % 100000}"

        if nid_val in seen_ids:
            errors.append(f"الصف {row_idx}: الرقم الوطني ({nid_val}) مكرر في الملف لطالب آخر.")

        seen_ids.add(nid_val)

        name_en_val = str(ws.cell(row=row_idx, column=col_mapping.get("full_name_en", 3)).value or "").strip()
        gender_raw = str(ws.cell(row=row_idx, column=col_mapping.get("gender", 4)).value or "").strip()
        gender = 2 if ("أنثى" in gender_raw or "female" in gender_raw.lower() or gender_raw == "2") else 1

        # Date of birth parsing
        dob_val = ws.cell(row=row_idx, column=col_mapping.get("date_of_birth", 5)).value
        dob_str = ""
        if isinstance(dob_val, datetime):
            dob_str = dob_val.strftime("%Y-%m-%d")
        elif dob_val:
            normalized_dob = normalize_digits(dob_val).replace("/", "-").replace(".", "-")
            match = re.search(r"\b(\d{4}-\d{1,2}-\d{1,2})\b", normalized_dob)
            if match:
                dob_str = match.group(1)
            else:
                dob_str = normalized_dob

        birthplace_val = str(ws.cell(row=row_idx, column=col_mapping.get("birthplace", 6)).value or "بغداد").strip()
        nationality_val = str(ws.cell(row=row_idx, column=col_mapping.get("nationality", 7)).value or "عراقي").strip()

        adm_year_raw = normalize_digits(ws.cell(row=row_idx, column=col_mapping.get("admission_year", 8)).value)
        adm_year = adm_year_raw if adm_year_raw else ""

        study_type_raw = str(ws.cell(row=row_idx, column=col_mapping.get("study_type", 9)).value or "").strip()
        study_type = "Evening" if "مسائي" in study_type_raw or "evening" in study_type_raw.lower() else "Morning"

        students.append({
            "national_id": nid_val,
            "full_name_ar": name_ar_val,
            "full_name_en": name_en_val,
            "gender": gender,
            "gender_label": "أنثى / Female" if gender == 2 else "ذكر / Male",
            "date_of_birth": dob_str,
            "birthplace": birthplace_val,
            "nationality": nationality_val,
            "admission_year": adm_year,
            "study_type": study_type,
            "study_type_ar": "مسائي" if study_type == "Evening" else "صباحي",
        })

    return students, errors


# -----------------------------------------------------------------------------
# Template 2: Courses & Grades Matrix
# -----------------------------------------------------------------------------

def generate_grades_matrix_template(
    students: List[Dict[str, Any]],
    courses: List[Dict[str, Any]],
    dept_name: str = "",
    stage_num: Optional[int] = None,
    academic_year: str = ""
) -> bytes:
    """
    Generate dynamic RTL Excel matrix template for grades.
    Pre-populates rows with students (if provided) and columns with department + shared courses.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "درجات المواد"
    ws.views.sheetView[0].rightToLeft = True

    total_cols = 2 + len(courses) + 3

    # Styles
    title_font = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    title_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")

    sub_font = Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
    sub_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    fixed_header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    fixed_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")

    dept_course_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    shared_course_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    course_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    attempt_fill = PatternFill(start_color="D97706", end_color="D97706", fill_type="solid")
    result_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    notes_fill = PatternFill(start_color="475569", end_color="475569", fill_type="solid")

    data_font = Font(name="Segoe UI", size=11)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Row 1: Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(total_cols, 6))
    title_cell = ws.cell(row=1, column=1)
    if stage_num:
        stage_arabic = {1: "الأولى", 2: "الثانية", 3: "الثالثة", 4: "الرابعة"}.get(stage_num, str(stage_num))
        title_cell.value = f"كشف درجات المواد — قسم: {dept_name or 'العام'} | المرحلة: {stage_arabic} | السنة: {academic_year or 'الحالية'}"
    else:
        title_cell.value = f"كشف درجات المواد — قسم: {dept_name or 'العام'} | السنة: {academic_year or 'الحالية'}"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = center_align
    ws.row_dimensions[1].height = 36

    # Row 2: Subtitle
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(total_cols, 6))
    sub_cell = ws.cell(row=2, column=1)
    sub_cell.value = "يرجى إدخال الدرجات بالأرقام من (0 إلى 100) في خانة كل مادة. الأعمدة الزرقاء: مواد تخصصية | الأعمدة الخضراء/الفيروزية: مواد مشتركة ومتطلبات كلية."
    sub_cell.font = sub_font
    sub_cell.fill = sub_fill
    sub_cell.alignment = center_align
    ws.row_dimensions[2].height = 24

    # Row 3: Headers
    ws.row_dimensions[3].height = 42

    # Fixed Col 1: National ID
    c1 = ws.cell(row=3, column=1)
    c1.value = "الرقم الوطني / الهوية"
    c1.font = fixed_header_font
    c1.fill = fixed_header_fill
    c1.alignment = center_align
    c1.border = thin_border
    ws.column_dimensions["A"].width = 24

    # Fixed Col 2: Name
    c2 = ws.cell(row=3, column=2)
    c2.value = "اسم الطالب (رباعي)"
    c2.font = fixed_header_font
    c2.fill = fixed_header_fill
    c2.alignment = center_align
    c2.border = thin_border
    ws.column_dimensions["B"].width = 34

    # Dynamic Course Columns
    current_col = 3
    for crs in courses:
        cell = ws.cell(row=3, column=current_col)
        crs_id = crs.get("id")
        crs_name = crs.get("name_ar") or crs.get("name_en") or f"مادة {crs_id}"
        units = crs.get("credit_hours", 3)
        crs_stage = int(crs.get("stage_number") or 1)
        is_shared = crs.get("is_shared", False)

        tag = " (مشترك)" if is_shared else ""
        cell.value = f"{crs_name}{tag}\n[ID:{crs_id}] (م{crs_stage} - {units} و)"
        cell.font = course_font
        cell.fill = shared_course_fill if is_shared else dept_course_fill
        cell.alignment = center_align
        cell.border = thin_border

        col_let = get_column_letter(current_col)
        ws.column_dimensions[col_let].width = max(22, len(crs_name) + 3)
        current_col += 1

    # Attempt Column
    c_att = ws.cell(row=3, column=current_col)
    c_att.value = "الدور\n(الأول/الثاني)"
    c_att.font = fixed_header_font
    c_att.fill = attempt_fill
    c_att.alignment = center_align
    c_att.border = thin_border
    ws.column_dimensions[get_column_letter(current_col)].width = 16
    current_col += 1

    # Result Column
    c_res = ws.cell(row=3, column=current_col)
    c_res.value = "النتيجة\n(ناجح/إكمال/راسب)"
    c_res.font = fixed_header_font
    c_res.fill = result_fill
    c_res.alignment = center_align
    c_res.border = thin_border
    ws.column_dimensions[get_column_letter(current_col)].width = 18
    current_col += 1

    # Notes Column
    c_not = ws.cell(row=3, column=current_col)
    c_not.value = "ملاحظات"
    c_not.font = fixed_header_font
    c_not.fill = notes_fill
    c_not.alignment = center_align
    c_not.border = thin_border
    ws.column_dimensions[get_column_letter(current_col)].width = 20

    # Populate Student Rows
    rows_to_render = students if students else [
        {"national_id": "199912345678", "full_name_ar": "علي حسن محمد جاسم"},
        {"national_id": "200087654321", "full_name_ar": "فاطمة أحمد حسين خضير"},
        {"national_id": "199854321098", "full_name_ar": "عمر خالد طارق سعيد"},
    ]

    for row_idx, st in enumerate(rows_to_render, start=4):
        ws.row_dimensions[row_idx].height = 26
        # Col 1: National ID
        c_nid = ws.cell(row=row_idx, column=1)
        c_nid.value = str(st.get("national_id") or "")
        c_nid.font = data_font
        c_nid.alignment = center_align
        c_nid.border = thin_border

        # Col 2: Name
        c_name = ws.cell(row=row_idx, column=2)
        c_name.value = str(st.get("full_name_ar") or "")
        c_name.font = data_font
        c_name.alignment = right_align
        c_name.border = thin_border

        # Empty Course cells
        c_idx = 3
        for _ in courses:
            c_score = ws.cell(row=row_idx, column=c_idx)
            c_score.font = data_font
            c_score.alignment = center_align
            c_score.border = thin_border
            c_idx += 1

        # Attempt default
        c_a = ws.cell(row=row_idx, column=c_idx)
        c_a.value = "الأول"
        c_a.font = data_font
        c_a.alignment = center_align
        c_a.border = thin_border
        c_idx += 1

        # Result default
        c_r = ws.cell(row=row_idx, column=c_idx)
        c_r.value = "ناجح"
        c_r.font = data_font
        c_r.alignment = center_align
        c_r.border = thin_border
        c_idx += 1

        # Notes
        c_n = ws.cell(row=row_idx, column=c_idx)
        c_n.font = data_font
        c_n.alignment = right_align
        c_n.border = thin_border

    # Dropdown validations
    att_col_letter = get_column_letter(2 + len(courses) + 1)
    res_col_letter = get_column_letter(2 + len(courses) + 2)

    dv_att = DataValidation(type="list", formula1='"الأول,الثاني,الثالث"', allow_blank=True)
    ws.add_data_validation(dv_att)
    dv_att.add(f"{att_col_letter}4:{att_col_letter}500")

    dv_res = DataValidation(type="list", formula1='"ناجح,إكمال,راسب,تأجيل"', allow_blank=True)
    ws.add_data_validation(dv_res)
    dv_res.add(f"{res_col_letter}4:{res_col_letter}500")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def parse_grades_file(
    file_bytes: bytes,
    catalog_courses: List[Dict[str, Any]]
) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Parse uploaded Grades Matrix Excel file.
    Matches course columns by [ID:xxx] or course title.
    Returns: (grades_list_per_student, error_messages)
    """
    grades_data: List[Dict[str, Any]] = []
    errors: List[str] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
    except Exception as exc:
        return [], [f"تعذر فتح ملف درجات الـ Excel: {exc}"]

    # Look for header row containing course IDs or column indicators
    header_row_idx = None
    for r in range(1, min(10, ws.max_row + 1)):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r, column=c).value or "").strip()
            if "ID:" in val or "اسم الطالب" in val or "الرقم الوطني" in val:
                header_row_idx = r
                break
        if header_row_idx:
            break

    if not header_row_idx:
        return [], ["لم يتم العثور على عناوين الدرجات والمواد في الملف. يرجى استخدام النموذج الرسمي."]

    # Map catalog courses by ID and by clean Arabic name
    course_by_id: Dict[int, Dict[str, Any]] = {int(c["id"]): c for c in catalog_courses if "id" in c}
    course_by_name: Dict[str, Dict[str, Any]] = {}
    for c in catalog_courses:
        n_ar = str(c.get("name_ar") or "").strip()
        if n_ar:
            course_by_name[n_ar] = c
        n_en = str(c.get("name_en") or "").strip()
        if n_en:
            course_by_name[n_en.lower()] = c

    # Detect column purposes
    col_to_course: Dict[int, Dict[str, Any]] = {}
    col_nid = 1
    col_name = 2
    col_round = None
    col_result = None

    for col in range(1, ws.max_column + 1):
        header_val = str(ws.cell(row=header_row_idx, column=col).value or "").strip()
        if not header_val:
            continue

        if "الرقم الوطني" in header_val or "الهوية" in header_val:
            col_nid = col
            continue
        if "اسم الطالب" in header_val:
            col_name = col
            continue
        if "الدور" in header_val:
            col_round = col
            continue
        if "النتيجة" in header_val:
            col_result = col
            continue

        # Match course by [ID:xxx]
        match_id = re.search(r"\[ID:(\d+)\]", header_val)
        if match_id:
            cid = int(match_id.group(1))
            if cid in course_by_id:
                col_to_course[col] = course_by_id[cid]
                continue

        # Match course by clean name
        clean_header = re.sub(r"\[ID:\d+\]", "", header_val)
        clean_header = re.sub(r"\(وحدات:[^\)]+\)", "", clean_header)
        clean_header = clean_header.replace("(مشترك)", "").replace("(تخصصي)", "").strip()

        if clean_header in course_by_name:
            col_to_course[col] = course_by_name[clean_header]
        elif clean_header.lower() in course_by_name:
            col_to_course[col] = course_by_name[clean_header.lower()]

    if not col_to_course:
        errors.append("لم يتم التعرف على أي أعمدة مواد دراسية مطابقة للمرحلة المختارة.")

    # Read row by row
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        name_val = str(ws.cell(row=row_idx, column=col_name).value or "").strip()
        nid_val = normalize_digits(ws.cell(row=row_idx, column=col_nid).value)

        if not name_val and not nid_val:
            continue  # empty row

        round_val = str(ws.cell(row=row_idx, column=col_round).value or "1") if col_round else "1"
        passed_round = 2 if "ثاني" in round_val or round_val == "2" else (3 if "ثالث" in round_val or round_val == "3" else 1)

        result_raw = str(ws.cell(row=row_idx, column=col_result).value or "ناجح") if col_result else "ناجح"
        if "راسب" in result_raw or "إعادة" in result_raw:
            result_status = "FAILED_REPEAT"
        elif "إكمال" in result_raw or "عبور" in result_raw:
            result_status = "CARRIED_OVER"
        elif "تأجيل" in result_raw:
            result_status = "DEFERRED"
        else:
            result_status = "PASSED"

        # Parse grades for each recognized course
        student_scores: List[Dict[str, Any]] = []
        weighted_sum = 0.0
        total_units = 0.0
        simple_sum = 0.0

        for col, crs in col_to_course.items():
            cell_val = ws.cell(row=row_idx, column=col).value
            score = parse_float_safe(cell_val)
            if score is None:
                continue

            if score < 0.0 or score > 100.0:
                errors.append(f"الصف {row_idx} ({name_val}): درجة المادة '{crs.get('name_ar')}' غير صحيحة ({score}). يجب أن تكون بين 0 و 100.")
                continue

            units = float(crs.get("credit_hours") or 3)
            weighted_sum += (score * units)
            total_units += units
            simple_sum += score

            student_scores.append({
                "course_id": crs["id"],
                "course_name_ar": crs.get("name_ar", ""),
                "course_name_en": crs.get("name_en", ""),
                "credit_hours": units,
                "score": score,
                "passed_round": passed_round,
                "stage_number": int(crs.get("stage_number") or 1),
            })

        computed_average = 0.0
        if total_units > 0:
            computed_average = round(weighted_sum / total_units, 3)
        elif student_scores:
            computed_average = round(simple_sum / len(student_scores), 3)

        grades_data.append({
            "national_id": nid_val,
            "student_name": name_val,
            "student_name_norm": " ".join(name_val.split()),
            "passed_round": passed_round,
            "result_status": result_status,
            "scores": student_scores,
            "scores_count": len(student_scores),
            "average": computed_average,
        })

    return grades_data, errors
